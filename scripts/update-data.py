from __future__ import annotations

import itertools
import json
import math
import re
import statistics
import time
import unicodedata
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
VALUES_PATH = ROOT / "indicator-values.js"
CATALOG_PATH = ROOT / "catalog-data.js"
PORDATA_SOURCES_PATH = ROOT / "data" / "pordata-sources.json"
PORDATA_DIRECTORY = ROOT / ".cache" / "pordata"
REPORT_PATH = ROOT / "data" / "update-report.json"

TERRITORIES = [
    "Alto Minho",
    "Arcos de Valdevez",
    "Caminha",
    "Melgaço",
    "Monção",
    "Paredes de Coura",
    "Ponte da Barca",
    "Ponte de Lima",
    "Valença",
    "Viana do Castelo",
    "Vila Nova de Cerveira",
]

INE_DIRECT = {
    "Superfície Km2": {"code": "0011866", "filters": {}},
    "Preço médio m²": {"code": "0012255", "filters": {"dim_3": "H1"}},
    "Rendimento médio": {"code": "0012649", "filters": {"dim_3": "TOT", "dim_4": "T"}},
    "Taxa de analfabetismo": {"code": "0012273", "filters": {"dim_3": "T"}},
    "Proporção de pop. residente em movimentos pendulares": {"code": "0011655", "filters": {}},
}


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", "".join(ch for ch in text if not unicodedata.combining(ch)).lower()).split())


NORMALIZED_TERRITORIES = {normalized(name): name for name in TERRITORIES}


def read_javascript_object(path: Path, variable: str) -> dict:
    source = path.read_text(encoding="utf-8")
    match = re.search(rf"const\s+{re.escape(variable)}\s*=\s*(.*);\s*$", source, re.DOTALL)
    if not match:
        raise RuntimeError(f"Não foi possível ler {variable} em {path.name}")
    return json.loads(match.group(1))


def write_javascript_object(path: Path, comment: str, variable: str, value: dict) -> None:
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    path.write_bytes(f"// {comment}\nconst {variable} = {payload};\n".encode("utf-8"))


def as_number(value: object) -> float | int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return int(value) if float(value).is_integer() else float(value)
    if isinstance(value, str):
        clean = value.strip().replace(" ", "").replace("%", "").replace("‰", "")
        if not clean or clean in {"-", "..", ":"}:
            return None
        if "," in clean and "." in clean:
            clean = clean.replace(".", "").replace(",", ".")
        elif "," in clean:
            clean = clean.replace(",", ".")
        try:
            number = float(clean)
            return int(number) if number.is_integer() else number
        except ValueError:
            return None
    return None


def validate_values(values: dict[str, float | int], indicator: str) -> None:
    missing = [name for name in TERRITORIES if name not in values or as_number(values[name]) is None]
    if missing:
        raise ValueError(f"{indicator}: faltam territórios: {', '.join(missing)}")
    if not any(abs(float(values[name])) > 0 for name in TERRITORIES):
        raise ValueError(f"{indicator}: todos os valores são zero")


def fetch_json(url: str, attempts: int = 3) -> object:
    request = urllib.request.Request(url, headers={"User-Agent": "Observatorio-Alto-Minho/1.0"})
    last_error = None
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                return json.load(response)
        except Exception as error:  # a rotina tenta novamente antes de falhar
            last_error = error
            time.sleep(2**attempt)
    raise RuntimeError(f"Falha ao consultar {url}: {last_error}")


def rows_for_latest_complete_year(item: dict, filters: dict[str, str]) -> tuple[str, list[dict]]:
    years = sorted((str(year) for year in item.get("Dados", {})), reverse=True)
    for year in years:
        rows = [
            row for row in item["Dados"][year]
            if normalized(row.get("geodsg")) in NORMALIZED_TERRITORIES
            and all(str(row.get(key)) == expected for key, expected in filters.items())
            and as_number(row.get("valor")) is not None
        ]
        if len({normalized(row["geodsg"]) for row in rows}) == len(TERRITORIES):
            return year, rows
    raise ValueError("O INE não devolveu um ano completo para os 11 territórios")


def update_ine(values_by_name: dict, report: list[dict]) -> int:
    updated = 0
    cache: dict[str, dict] = {}
    for name, config in INE_DIRECT.items():
        code = config["code"]
        if code not in cache:
            payload = fetch_json(f"https://www.ine.pt/ine/json_indicador/pindica.jsp?op=2&varcd={code}&lang=PT")
            if not isinstance(payload, list) or not payload or not payload[0].get("Sucesso", {}).get("Verdadeiro"):
                raise ValueError(f"INE {code}: resposta inválida")
            cache[code] = payload[0]
        year, rows = rows_for_latest_complete_year(cache[code], config["filters"])
        current_year = int(values_by_name[name]["year"])
        if int(year) < current_year:
            report.append({"indicator": name, "source": "INE", "status": "kept", "reason": "source year is older"})
            continue
        new_values = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: as_number(row["valor"]) for row in rows}
        validate_values(new_values, name)
        changed = values_by_name[name]["values"] != new_values or values_by_name[name]["year"] != year
        values_by_name[name]["values"] = new_values
        values_by_name[name]["year"] = year
        updated += 1
        report.append({"indicator": name, "source": "INE", "status": "updated" if changed else "checked", "year": year})

    name = "Diferença salarial H/M (%)"
    item = cache.get("0012649")
    if item is None:
        payload = fetch_json("https://www.ine.pt/ine/json_indicador/pindica.jsp?op=2&varcd=0012649&lang=PT")
        item = payload[0]
    year_h, men_rows = rows_for_latest_complete_year(item, {"dim_3": "TOT", "dim_4": "1"})
    year_m, women_rows = rows_for_latest_complete_year(item, {"dim_3": "TOT", "dim_4": "2"})
    if year_h != year_m:
        raise ValueError("INE 0012649: os anos por sexo não coincidem")
    men = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: float(row["valor"]) for row in men_rows}
    women = {NORMALIZED_TERRITORIES[normalized(row["geodsg"])]: float(row["valor"]) for row in women_rows}
    gap = {territory: round((men[territory] - women[territory]) / men[territory] * 100, 4) for territory in TERRITORIES}
    validate_values(gap, name)
    if int(year_h) >= int(values_by_name[name]["year"]):
        changed = values_by_name[name]["values"] != gap or values_by_name[name]["year"] != year_h
        values_by_name[name]["values"] = gap
        values_by_name[name]["year"] = year_h
        updated += 1
        report.append({"indicator": name, "source": "INE", "status": "updated" if changed else "checked", "year": year_h})
    return updated


def pordata_candidates(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    territory_rows: dict[str, int] = {}
    for row_index, row in enumerate(rows):
        if len(row) < 2:
            continue
        territory = NORMALIZED_TERRITORIES.get(normalized(row[1]))
        if territory:
            territory_rows.setdefault(territory, row_index)
    if len(territory_rows) != len(TERRITORIES):
        raise ValueError(f"{path.name}: só foram encontrados {len(territory_rows)} territórios")
    year_header_row = next(
        (
            row_index for row_index, row in enumerate(rows)
            if row and normalized(row[0]) == "ambito geografico" and len(row) > 1 and normalized(row[1]) == "anos"
        ),
        None,
    )
    if year_header_row is None:
        raise ValueError(f"{path.name}: cabeçalho de anos não encontrado")
    max_columns = max(len(row) for row in rows)
    candidates = []
    for column in range(2, max_columns):
        year = as_number(rows[year_header_row][column]) if column < len(rows[year_header_row]) else None
        if not isinstance(year, int) or not 1900 <= year <= 2100:
            year = None
        if year is None:
            continue
        column_values = {}
        for territory, row_index in territory_rows.items():
            cell = rows[row_index][column] if column < len(rows[row_index]) else None
            number = as_number(cell)
            if number is not None:
                column_values[territory] = number
        if len(column_values) == len(TERRITORIES):
            candidates.append({"column": column, "year": year, "values": column_values})
    workbook.close()
    return candidates


def candidate_score(candidate: dict[str, float | int], current: dict[str, float | int]) -> float:
    relative_errors = []
    for territory in TERRITORIES:
        old = float(current[territory])
        new = float(candidate[territory])
        relative_errors.append(abs(new - old) / max(abs(old), 1.0))
    return statistics.median(relative_errors)


def choose_pordata_values(path: Path, current: dict, old_year: int) -> tuple[str, dict, float, tuple[int, ...]]:
    candidates = pordata_candidates(path)
    if not candidates:
        raise ValueError(f"{path.name}: não há colunas de dados completas")
    newest_year = max(candidate["year"] for candidate in candidates)
    if newest_year < old_year:
        raise ValueError(f"{path.name}: o ano {newest_year} é anterior a {old_year}")
    latest = [candidate for candidate in candidates if candidate["year"] == newest_year]
    options = []
    maximum_size = min(4, len(latest))
    for size in range(1, maximum_size + 1):
        for combination in itertools.combinations(latest, size):
            combined = {territory: sum(float(part["values"][territory]) for part in combination) for territory in TERRITORIES}
            combined = {key: int(value) if value.is_integer() else value for key, value in combined.items()}
            options.append((candidate_score(combined, current), combined, tuple(part["column"] for part in combination)))
            if len(options) >= 2000:
                break
        if len(options) >= 2000:
            break
    score, selected, columns = min(options, key=lambda option: option[0])
    if score > 0.80:
        raise ValueError(f"{path.name}: nenhuma série é suficientemente próxima (erro mediano {score:.2f})")
    validate_values(selected, path.name)
    return str(newest_year), selected, score, columns


def update_pordata(values_by_name: dict, report: list[dict]) -> int:
    sources = json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8"))
    updated = 0
    for name, url in sources.items():
        if name == "Diferença entre o salário mínimo nacional e a remuneração base média mensal dos trabalhadores por conta de outrem":
            continue
        source_id = re.search(r"-(\d+)$", url).group(1)
        path = PORDATA_DIRECTORY / f"{source_id}.xlsx"
        if not path.exists():
            raise FileNotFoundError(f"Falta o ficheiro Pordata {path.name}")
        current = values_by_name[name]
        try:
            year, selected, score, columns = choose_pordata_values(path, current["values"], int(current["year"]))
            changed = current["values"] != selected or current["year"] != year
            current["values"] = selected
            current["year"] = year
            updated += 1
            report.append({"indicator": name, "source": "Pordata", "status": "updated" if changed else "checked", "year": year, "match_error": round(score, 4), "columns": columns})
        except Exception as error:
            report.append({"indicator": name, "source": "Pordata", "status": "kept", "reason": str(error)})
    alias = "Pop. Estrangeira c/ estatuto residente"
    source = "População estrangeira com estatuto legal de residente: total e por algumas nacionalidades"
    values_by_name[alias]["year"] = values_by_name[source]["year"]
    values_by_name[alias]["values"] = dict(values_by_name[source]["values"])
    return updated


def update_salary_difference(values_by_name: dict, report: list[dict]) -> int:
    name = "Diferença entre o salário mínimo nacional e a remuneração base média mensal dos trabalhadores por conta de outrem"
    income = values_by_name["Rendimento médio"]
    year = int(income["year"])
    workbook_path = PORDATA_DIRECTORY / "74.xlsx"
    if not workbook_path.exists():
        raise FileNotFoundError("Falta o ficheiro Pordata 74.xlsx com o salário mínimo")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    minimum_wage = None
    for row in workbook.worksheets[0].iter_rows(values_only=True):
        if row and as_number(row[0]) == year and len(row) > 2:
            minimum_wage = as_number(row[2])
            break
    workbook.close()
    if minimum_wage is None or minimum_wage <= 0:
        raise ValueError(f"Pordata 74: salário mínimo indisponível para {year}")
    difference = {
        territory: round(float(income["values"][territory]) - float(minimum_wage), 1)
        for territory in TERRITORIES
    }
    validate_values(difference, name)
    changed = values_by_name[name]["values"] != difference or values_by_name[name]["year"] != str(year)
    values_by_name[name]["values"] = difference
    values_by_name[name]["year"] = str(year)
    report.append({
        "indicator": name,
        "source": "Pordata + INE",
        "status": "updated" if changed else "checked",
        "year": str(year),
        "minimum_wage": minimum_wage,
    })
    return 1


def update_catalog_links() -> None:
    catalog = read_javascript_object(CATALOG_PATH, "indicatorCatalog")
    sources = json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8"))
    for item in catalog:
        if item["name"] in sources:
            item["url"] = sources[item["name"]]
        if item["name"] in INE_DIRECT:
            code = INE_DIRECT[item["name"]]["code"]
            item["url"] = f"https://www.ine.pt/xportal/xmain?xpid=INE&xpgid=ine_indicadores&indOcorrCod={code}&contexto=bd&selTab=tab2"
        if item["name"] == "Diferença salarial H/M (%)":
            item["url"] = "https://www.ine.pt/xportal/xmain?xpid=INE&xpgid=ine_indicadores&indOcorrCod=0012649&contexto=bd&selTab=tab2"
        if item["name"] == "Pop. Estrangeira c/ estatuto residente":
            item["source"] = "Pordata"
            item["url"] = sources["População estrangeira com estatuto legal de residente: total e por algumas nacionalidades"]
    write_javascript_object(CATALOG_PATH, "Catálogo da matriz intermunicipal e ligações permanentes das fontes.", "indicatorCatalog", catalog)


def main() -> None:
    values_by_name = read_javascript_object(VALUES_PATH, "indicatorValuesByName")
    expected_pordata = len(json.loads(PORDATA_SOURCES_PATH.read_text(encoding="utf-8")))
    report: list[dict] = []
    pordata_count = update_pordata(values_by_name, report)
    ine_count = update_ine(values_by_name, report)
    pordata_count += update_salary_difference(values_by_name, report)
    if pordata_count != expected_pordata:
        failures = [entry for entry in report if entry["source"] == "Pordata" and entry["status"] == "kept"]
        raise RuntimeError(f"Só {pordata_count} de {expected_pordata} indicadores Pordata passaram a validação. Falhas: {failures}")
    if ine_count != 6:
        raise RuntimeError(f"Só {ine_count} indicadores INE passaram a validação")
    write_javascript_object(VALUES_PATH, "Valores atualizados automaticamente a partir do INE e da Pordata.", "indicatorValuesByName", values_by_name)
    update_catalog_links()
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Atualização concluída: {pordata_count} indicadores Pordata e {ine_count} indicadores INE verificados.")


if __name__ == "__main__":
    main()
